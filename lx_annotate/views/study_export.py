from __future__ import annotations

from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Final, Literal, TypedDict, cast

from django.db.models.expressions import Exists, OuterRef
from django.db.models.query import Prefetch, QuerySet
from django.db.models.query_utils import Q
from django.http import HttpResponse
from django.utils.timezone import localdate, now
from endoreg_db.authz.permissions import satisfies
from endoreg_db.models import (
    PatientExamination,
    PatientExaminationIndication,
    PatientExaminationReport,
    PatientFinding,
    RawPdfFile,
    VideoFile,
)
from endoreg_db.services.study_cohort import (
    StudyCohortFilters,
    parse_study_cohort_filters,
)
from endoreg_db.utils.permissions import EnvironmentAwarePermission, is_debug_mode
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import BasePermission
from rest_framework.request import Request
from rest_framework.response import Response

MAX_STUDY_EXPORT_ROWS: Final = 25_000
FORMULA_PREFIXES: Final = ("=", "+", "-", "@")
type StudyExportGrouping = Literal["patient", "examination"]
type StudyExportMode = Literal["criteria", "cohort"]


class StudyExportOptions(TypedDict):
    examinations: list[str]
    findings: list[str]
    indications: list[str]
    maximum_rows: int


@dataclass(frozen=True)
class StudyExportSelection:
    mode: StudyExportMode = "criteria"
    examinations: tuple[str, ...] = ()
    findings: tuple[str, ...] = ()
    indications: tuple[str, ...] = ()
    group_by: StudyExportGrouping = "patient"
    study_name: str = ""
    hypothesis: str = ""
    cohort_schema_version: str = ""
    patient_examination_ids: tuple[int, ...] = ()
    cohort_filters: StudyCohortFilters | None = None

    @property
    def has_criteria(self) -> bool:
        return bool(self.examinations or self.findings or self.indications)


@dataclass(slots=True)
class _GroupedStudyRow:
    patient_hash: str
    center_key: str
    patient_examination_ids: list[int]
    case_hashes: list[str]
    examination_dates: list[str]
    examination_names: list[str]
    indications: set[str]
    findings: set[str]
    report_ids: set[int]
    report_file_ids: set[int]
    video_ids: set[int]


class StudyExportPatientReadPermission(BasePermission):
    """Apply the same patient-read role used by the cohort preview endpoint."""

    def has_permission(  # pyright: ignore[reportIncompatibleMethodOverride]
        self, request: Request, view: object
    ) -> bool:
        del view
        if is_debug_mode():
            return True
        user = request.user
        if not getattr(user, "is_authenticated", False):
            return False
        groups = getattr(user, "groups", None)
        if groups is None:
            return False
        roles = set(groups.values_list("name", flat=True))
        return satisfies(roles, "patient:read")


def _pseudonymous_cases() -> QuerySet[PatientExamination]:
    return (
        PatientExamination.objects.select_related(
            "patient",
            "patient__center",
            "examination",
        )
        .filter(
            patient__is_real_person=False,
            patient__patient_hash__isnull=False,
        )
        .exclude(patient__patient_hash="")
    )


def _sorted_strings(values: Iterable[object]) -> list[str]:
    return sorted({str(value).strip() for value in values if str(value).strip()})


def build_study_export_options() -> StudyExportOptions:
    cases = _pseudonymous_cases()
    case_ids = cases.values("pk")
    examinations = cases.values_list("examination__name", flat=True)
    findings = PatientFinding.objects.filter(
        patient_examination_id__in=case_ids,
        is_active=True,
    ).values_list("finding__name", flat=True)
    indications = PatientExaminationIndication.objects.filter(
        patient_examination_id__in=case_ids,
    ).values_list("examination_indication__name", flat=True)
    return {
        "examinations": _sorted_strings(examinations),
        "findings": _sorted_strings(findings),
        "indications": _sorted_strings(indications),
        "maximum_rows": MAX_STUDY_EXPORT_ROWS,
    }


def _query_values(query_params: Mapping[str, Any], key: str) -> tuple[str, ...]:
    getlist = getattr(query_params, "getlist", None)
    raw_value = query_params.get(key, "")
    if callable(getlist):
        raw_values = cast(Iterable[Any], getlist(key))
    elif isinstance(raw_value, (list, tuple)):
        raw_values = raw_value
    else:
        raw_values = [raw_value]
    values: list[str] = []
    for raw_value in raw_values:
        value = str(raw_value or "").strip()
        if value and value not in values:
            values.append(value)
    return tuple(values)


def _query_positive_ints(query_params: Mapping[str, Any], key: str) -> tuple[int, ...]:
    values = _query_values(query_params, key)
    parsed: list[int] = []
    for value in values:
        try:
            parsed_value = int(value)
        except ValueError as exc:
            raise ValueError(f"{key} must contain positive integers.") from exc
        if parsed_value < 1:
            raise ValueError(f"{key} must contain positive integers.")
        if parsed_value not in parsed:
            parsed.append(parsed_value)
    return tuple(parsed)


def _cohort_filter_mapping(query_params: Mapping[str, Any]) -> Mapping[str, Any]:
    normalized = dict(query_params)
    for key in ("has_report", "has_video"):
        value = normalized.get(key)
        if isinstance(value, bool):
            normalized[key] = "true" if value else "false"
    return normalized


def parse_study_export_selection(
    query_params: Mapping[str, Any],
) -> StudyExportSelection:
    mode = str(query_params.get("mode", "criteria") or "criteria").strip()
    if mode not in {"criteria", "cohort"}:
        raise ValueError("mode must be criteria or cohort.")
    if mode == "cohort":
        if _parse_grouping(query_params) != "patient":
            raise ValueError("Study cohort exports must use patient grouping.")
        study_name = str(query_params.get("study_name", "") or "").strip()
        hypothesis = str(query_params.get("hypothesis", "") or "").strip()
        patient_examination_ids = _query_positive_ints(
            query_params, "patient_examination_id"
        )
        if not study_name or not hypothesis:
            raise ValueError("Study cohort exports require study_name and hypothesis.")
        if not patient_examination_ids:
            raise ValueError(
                "Study cohort exports require at least one patient_examination_id."
            )
        return StudyExportSelection(
            mode="cohort",
            group_by="patient",
            study_name=study_name,
            hypothesis=hypothesis,
            cohort_schema_version=str(
                query_params.get("cohort_schema_version", "") or ""
            ).strip(),
            patient_examination_ids=patient_examination_ids,
            cohort_filters=parse_study_cohort_filters(
                _cohort_filter_mapping(query_params)
            ),
        )

    selection = StudyExportSelection(
        mode="criteria",
        examinations=_query_values(query_params, "examination"),
        findings=_query_values(query_params, "finding"),
        indications=_query_values(query_params, "indication"),
        group_by=_parse_grouping(query_params),
    )
    if not selection.has_criteria:
        raise ValueError("Select at least one examination, finding, or indication.")

    options = build_study_export_options()
    for label, selected, available in (
        ("examination", selection.examinations, options["examinations"]),
        ("finding", selection.findings, options["findings"]),
        ("indication", selection.indications, options["indications"]),
    ):
        unknown = sorted(set(selected) - set(available))
        if unknown:
            raise ValueError(f"Unknown {label} value(s): {', '.join(unknown)}")
    return selection


def _parse_grouping(query_params: Mapping[str, Any]) -> StudyExportGrouping:
    value = str(query_params.get("group_by", "patient") or "patient").strip()
    if value not in {"patient", "examination"}:
        raise ValueError("group_by must be patient or examination.")
    return cast(StudyExportGrouping, value)


def _matching_cases(
    selection: StudyExportSelection,
) -> QuerySet[PatientExamination]:
    cases = _pseudonymous_cases()
    if selection.examinations:
        cases = cases.filter(examination__name__in=selection.examinations)
    if selection.findings:
        cases = cases.filter(
            patient_findings__is_active=True,
            patient_findings__finding__name__in=selection.findings,
        )
    if selection.indications:
        cases = cases.filter(
            indications__examination_indication__name__in=selection.indications,
        )
    return cases.distinct()


def _export_cases(
    selection: StudyExportSelection,
) -> tuple[QuerySet[PatientExamination], int]:
    matching_cases = _matching_cases(selection)
    active_findings = PatientFinding.objects.filter(is_active=True).select_related(
        "finding"
    )
    indications = PatientExaminationIndication.objects.select_related(
        "examination_indication"
    )
    active_reports = PatientExaminationReport.objects.filter(is_active=True).only(
        "pk", "patient_examination_id"
    )
    eligible_report_files = (
        RawPdfFile.objects.filter(
            state__anonymization_validated=True,
            state__processed_file_sha256__gt="",
        )
        .exclude(processed_file="")
        .only("pk", "examination_id")
    )
    eligible_videos = (
        VideoFile.objects.filter(
            state__anonymization_validated=True,
            state__processed_file_sha256__gt="",
        )
        .exclude(processed_file="")
        .only("pk", "examination_id")
    )
    prefetches = (
        Prefetch("patient_findings", queryset=active_findings),
        Prefetch("indications", queryset=indications),
        Prefetch("reports", queryset=active_reports),
        Prefetch("raw_pdf_files", queryset=eligible_report_files),
        Prefetch("video_files", queryset=eligible_videos),
    )
    if selection.mode == "cohort":
        cases = _pseudonymous_cases().filter(pk__in=selection.patient_examination_ids)
        if cases.count() != len(selection.patient_examination_ids):
            raise ValueError(
                "The study cohort contains an unavailable patient examination."
            )
        row_count = (
            cases.values("patient__center_id", "patient__patient_hash")
            .distinct()
            .count()
        )
        return (
            cases.prefetch_related(*prefetches).order_by(
                "patient__center__center_key",
                "patient__patient_hash",
                "date_start",
                "pk",
            ),
            row_count,
        )
    if selection.group_by == "examination":
        return (
            matching_cases.prefetch_related(*prefetches).order_by(
                "date_start",
                "pk",
            ),
            matching_cases.count(),
        )

    row_count = (
        matching_cases.values("patient__center_id", "patient__patient_hash")
        .distinct()
        .count()
    )
    scoped_match = matching_cases.order_by().filter(
        patient__center_id=OuterRef("patient__center_id"),
        patient__patient_hash=OuterRef("patient__patient_hash"),
    )
    unscoped_match = matching_cases.order_by().filter(
        patient__center_id__isnull=True,
        patient__patient_hash=OuterRef("patient__patient_hash"),
    )
    cases = (
        _pseudonymous_cases()
        .annotate(
            scoped_identity_match=Exists(scoped_match),
            unscoped_identity_match=Exists(unscoped_match),
        )
        .filter(
            Q(patient__center_id__isnull=False, scoped_identity_match=True)
            | Q(patient__center_id__isnull=True, unscoped_identity_match=True)
        )
        .prefetch_related(*prefetches)
        .order_by(
            "patient__center__center_key",
            "patient__patient_hash",
            "date_start",
            "pk",
        )
    )
    return cases, row_count


def _safe_cell(value: object | None) -> str:
    text = str(value or "").strip()
    if text.startswith(FORMULA_PREFIXES):
        return f"'{text}"
    return text


def _joined(values: Iterable[object]) -> str:
    return "; ".join(_sorted_strings(values))


def _ordered_join(values: Iterable[object]) -> str:
    return "; ".join(str(value).strip() for value in values if str(value).strip())


def _group_key(case: PatientExamination) -> tuple[str, str]:
    center = case.patient.center
    center_key = str(getattr(center, "center_key", "") or "").strip()
    patient_hash = str(case.patient.patient_hash or "").strip()
    return center_key, patient_hash


def _group_cases(
    cases: Iterable[PatientExamination],
) -> list[_GroupedStudyRow]:
    grouped: dict[tuple[str, str], _GroupedStudyRow] = {}
    for case in cases:
        center_key, patient_hash = _group_key(case)
        row = grouped.setdefault(
            (center_key, patient_hash),
            _GroupedStudyRow(
                patient_hash=patient_hash,
                center_key=center_key,
                patient_examination_ids=[],
                case_hashes=[],
                examination_dates=[],
                examination_names=[],
                indications=set(),
                findings=set(),
                report_ids=set(),
                report_file_ids=set(),
                video_ids=set(),
            ),
        )
        row.patient_examination_ids.append(case.pk)
        row.case_hashes.append(str(case.hash))
        row.examination_dates.append(
            case.date_start.isoformat() if case.date_start else ""
        )
        row.examination_names.append(str(getattr(case.examination, "name", "") or ""))
        row.findings.update(
            str(finding.finding.name) for finding in case.patient_findings.all()
        )
        row.indications.update(
            str(indication.examination_indication.name)
            for indication in case.indications.all()
        )
        row.report_ids.update(report.pk for report in case.reports.all())
        row.report_file_ids.update(report.pk for report in case.raw_pdf_files.all())
        row.video_ids.update(video.pk for video in case.video_files.all())
    return list(grouped.values())


def _styled_header_cell(worksheet: Any, value: str) -> Any:
    cell = WriteOnlyCell(worksheet, value=value)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    return cell


def _selection_rows(selection: StudyExportSelection) -> list[tuple[str, str]]:
    if selection.mode == "cohort":
        filters = selection.cohort_filters or StudyCohortFilters()
        return [
            ("Export mode", "cohort"),
            ("Grouping", selection.group_by),
            ("Study name", selection.study_name),
            ("Hypothesis", selection.hypothesis),
            ("Cohort schema version", selection.cohort_schema_version),
            ("Date from", filters.date_from.isoformat() if filters.date_from else ""),
            ("Date to", filters.date_to.isoformat() if filters.date_to else ""),
            ("Center", filters.center_key),
            ("Examination", filters.examination_name),
            ("Document type", filters.document_type),
            ("Finding", filters.finding),
            ("Annotation label", filters.annotation_label),
            (
                "Has report",
                str(filters.has_report) if filters.has_report is not None else "",
            ),
            (
                "Has video",
                str(filters.has_video) if filters.has_video is not None else "",
            ),
            ("Preview limit", str(filters.limit)),
        ]
    return [
        ("Export mode", "criteria"),
        ("Grouping", selection.group_by),
        ("Examinations", _joined(selection.examinations)),
        ("Findings", _joined(selection.findings)),
        ("Indications", _joined(selection.indications)),
    ]


def _examination_values(case: PatientExamination) -> list[str]:
    center_key, patient_hash = _group_key(case)
    return [
        _safe_cell(case.hash),
        _safe_cell(patient_hash),
        _safe_cell(center_key),
        str(case.pk),
        case.date_start.isoformat() if case.date_start else "",
        _safe_cell(getattr(case.examination, "name", "")),
        _safe_cell(
            _joined(
                indication.examination_indication.name
                for indication in case.indications.all()
            )
        ),
        _safe_cell(
            _joined(finding.finding.name for finding in case.patient_findings.all())
        ),
        _safe_cell(_joined(report.pk for report in case.reports.all())),
        _safe_cell(_joined(report.pk for report in case.raw_pdf_files.all())),
        _safe_cell(_joined(video.pk for video in case.video_files.all())),
    ]


def build_study_export_workbook(
    cases: Iterable[PatientExamination],
    *,
    selection: StudyExportSelection,
    row_count: int,
) -> bytes:
    workbook = Workbook(write_only=True)
    cases_sheet = workbook.create_sheet("Cases")
    headers = (
        (
            "Patient pseudonym",
            "Center",
            "Patient examination IDs",
            "Case pseudonyms",
            "Examination dates",
            "Examinations",
            "Indications",
            "Findings",
            "Report IDs",
            "Processed report file IDs",
            "Processed video IDs",
        )
        if selection.group_by == "patient"
        else (
            "Case pseudonym",
            "Patient pseudonym",
            "Center",
            "Patient examination ID",
            "Examination date",
            "Examination",
            "Indications",
            "Findings",
            "Report IDs",
            "Processed report file IDs",
            "Processed video IDs",
        )
    )
    cases_sheet.append([_styled_header_cell(cases_sheet, value) for value in headers])
    cases_sheet.freeze_panes = "A2"

    if selection.group_by == "patient":
        grouped_rows = _group_cases(cases)
        if len(grouped_rows) != row_count:
            raise RuntimeError(
                "Pseudonymous study export grouping changed while the workbook was built."
            )
        for row in grouped_rows:
            cases_sheet.append(
                [
                    _safe_cell(row.patient_hash),
                    _safe_cell(row.center_key),
                    _safe_cell(_ordered_join(row.patient_examination_ids)),
                    _safe_cell(_ordered_join(row.case_hashes)),
                    _safe_cell(_ordered_join(row.examination_dates)),
                    _safe_cell(_ordered_join(row.examination_names)),
                    _safe_cell(_joined(row.indications)),
                    _safe_cell(_joined(row.findings)),
                    _safe_cell(_joined(row.report_ids)),
                    _safe_cell(_joined(row.report_file_ids)),
                    _safe_cell(_joined(row.video_ids)),
                ]
            )
    else:
        for case in cases:
            cases_sheet.append(_examination_values(case))
    cases_sheet.auto_filter.ref = f"A1:K{row_count + 1}"

    metadata_sheet = workbook.create_sheet("Export metadata")
    metadata_sheet.append(
        [
            _styled_header_cell(metadata_sheet, "Property"),
            _styled_header_cell(metadata_sheet, "Value"),
        ]
    )
    metadata_sheet.append(["Schema version", "1.0"])
    metadata_sheet.append(["Pseudonymous", "Yes"])
    metadata_sheet.append(["Row count", row_count])
    filter_semantics = (
        "Exact reviewed StudyCohortPage snapshot"
        if selection.mode == "cohort"
        else "OR within categories; AND across categories"
    )
    metadata_sheet.append(["Filter semantics", filter_semantics])
    for label, values in _selection_rows(selection):
        metadata_sheet.append([label, _safe_cell(values)])
    metadata_sheet.append(["Generated at (UTC)", now().isoformat()])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _selection_from_request(request: Request) -> StudyExportSelection:
    request_values = cast(
        object,
        request.data if request.method == "POST" else request.query_params,
    )
    mapping = (
        cast(Mapping[str, Any], request_values)
        if isinstance(request_values, Mapping)
        else {}
    )
    return parse_study_export_selection(mapping)


@api_view(["GET"])
@permission_classes([EnvironmentAwarePermission, StudyExportPatientReadPermission])
def study_export_options(request: Request) -> Response:
    del request
    return Response(build_study_export_options(), status=status.HTTP_200_OK)


@api_view(["GET", "POST"])
@permission_classes([EnvironmentAwarePermission, StudyExportPatientReadPermission])
def study_export_xlsx(request: Request) -> Response | HttpResponse:
    try:
        selection = _selection_from_request(request)
        cases, row_count = _export_cases(selection)
    except ValueError as exc:
        return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    if row_count == 0:
        return Response(
            {"error": "No pseudonymous cases match the selected criteria."},
            status=status.HTTP_404_NOT_FOUND,
        )
    if row_count > MAX_STUDY_EXPORT_ROWS:
        return Response(
            {
                "error": (
                    f"The export contains {row_count} rows; narrow the selection to "
                    f"at most {MAX_STUDY_EXPORT_ROWS} rows."
                )
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    case_rows = cast(
        Iterable[PatientExamination],
        cases.iterator(chunk_size=500),
    )
    workbook = build_study_export_workbook(
        case_rows,
        selection=selection,
        row_count=row_count,
    )
    export_label = "cohort" if selection.mode == "cohort" else "cases"
    filename = f"pseudonymous-study-{export_label}-{localdate().isoformat()}.xlsx"
    response = HttpResponse(
        workbook,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["X-Export-Row-Count"] = str(row_count)
    response["X-Content-Type-Options"] = "nosniff"
    return response
