from __future__ import annotations

from collections.abc import Iterator
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from typing import Any, cast

import pytest
from django.test import Client, override_settings
from endoreg_db.models import (
    Center,
    Examination,
    ExaminationIndication,
    Finding,
    Patient,
    PatientExamination,
    PatientExaminationIndication,
    PatientExaminationReport,
    PatientFinding,
)
from openpyxl import load_workbook
from rest_framework.request import Request

from lx_annotate.views.study_export import StudyExportPatientReadPermission

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _study_export_settings() -> Iterator[None]:
    with override_settings(
        ROOT_URLCONF="lx_annotate.urls",
        ALLOWED_HOSTS=["testserver"],
        DEBUG=True,
    ):
        yield


def _patient(
    *,
    patient_hash: str,
    center: Center,
    real: bool = False,
) -> Patient:
    return Patient.objects.create(
        first_name="Sensitive",
        last_name="Identity",
        dob=date(1970, 1, 1),
        patient_hash=patient_hash,
        center=center,
        is_real_person=real,
    )


def _case(
    *,
    patient: Patient,
    examination: Examination,
    case_hash: str,
    finding: Finding | None = None,
    indication: ExaminationIndication | None = None,
    examination_date: date = date(2026, 8, 1),
) -> PatientExamination:
    case = PatientExamination.objects.create(
        patient=patient,
        examination=examination,
        date_start=examination_date,
        hash=case_hash,
    )
    if finding is not None:
        PatientFinding.objects.create(patient_examination=case, finding=finding)
    if indication is not None:
        PatientExaminationIndication.objects.create(
            patient_examination=case,
            examination_indication=indication,
        )
    return case


@pytest.fixture
def annotated_cases() -> dict[str, object]:
    center_a = Center.objects.create(name="Center A", center_key="center-a")
    center_b = Center.objects.create(name="Center B", center_key="center-b")
    colonoscopy = Examination.objects.create(name="colonoscopy")
    gastroscopy = Examination.objects.create(name="gastroscopy")
    polyp = Finding.objects.create(name="polyp")
    gastritis = Finding.objects.create(name="gastritis")
    screening = ExaminationIndication.objects.create(name="screening")
    surveillance = ExaminationIndication.objects.create(name="surveillance")

    matching_case = _case(
        patient=_patient(patient_hash="patient-pseudo-1", center=center_a),
        examination=colonoscopy,
        case_hash="=case-formula-prefix",
        finding=polyp,
        indication=screening,
    )
    follow_up_case = _case(
        patient=_patient(patient_hash="patient-pseudo-1", center=center_a),
        examination=gastroscopy,
        case_hash="case-follow-up",
        finding=gastritis,
        indication=surveillance,
        examination_date=date(2026, 9, 1),
    )
    follow_up_report = PatientExaminationReport.objects.create(
        patient_examination=follow_up_case,
        template_name="follow-up-report",
    )
    _case(
        patient=_patient(patient_hash="patient-pseudo-2", center=center_a),
        examination=colonoscopy,
        case_hash="case-2",
        finding=gastritis,
        indication=screening,
    )
    _case(
        patient=_patient(patient_hash="patient-pseudo-3", center=center_a),
        examination=gastroscopy,
        case_hash="case-3",
        finding=polyp,
        indication=surveillance,
    )
    _case(
        patient=_patient(patient_hash="real-patient", center=center_a, real=True),
        examination=colonoscopy,
        case_hash="real-case",
        finding=polyp,
        indication=screening,
    )
    _case(
        patient=_patient(patient_hash="patient-pseudo-1", center=center_b),
        examination=gastroscopy,
        case_hash="other-center-case",
        finding=gastritis,
        indication=surveillance,
        examination_date=date(2026, 10, 1),
    )
    return {
        "matching_case": matching_case,
        "follow_up_case": follow_up_case,
        "follow_up_report": follow_up_report,
    }


def test_options_only_include_concepts_from_pseudonymous_cases(
    annotated_cases: dict[str, object],
) -> None:
    del annotated_cases
    response = Client().get("/endoreg-api/media/studies/case-export/options/")

    assert response.status_code == 200
    assert response.json() == {
        "examinations": ["colonoscopy", "gastroscopy"],
        "findings": ["gastritis", "polyp"],
        "indications": ["screening", "surveillance"],
        "maximum_rows": 25_000,
    }


def test_xlsx_uses_or_within_categories_and_and_across_categories(
    annotated_cases: dict[str, object],
) -> None:
    # Arrange
    request_query = [
        ("examination", "colonoscopy"),
        ("examination", "gastroscopy"),
        ("finding", "polyp"),
        ("indication", "screening"),
    ]

    # Act
    response = Client().get(
        "/endoreg-api/media/studies/case-export.xlsx",
        request_query,
        HTTP_ACCEPT="application/json",
    )

    # Assert
    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response["X-Export-Row-Count"] == "1"
    assert response["Content-Disposition"].endswith('.xlsx"')

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    cases = list(workbook["Cases"].iter_rows(values_only=True))
    metadata = dict(workbook["Export metadata"].iter_rows(values_only=True))

    assert cases[0] == (
        "Patient pseudonym",
        "Center",
        "Patient examination IDs",
        "Case pseudonyms",
        "Examination dates",
        "Examinations",
        "Indications",
        "Findings",
        "Report IDs",
        "Processed report file IDs",
        "Processed video IDs",
    )
    matching_case = cast(PatientExamination, annotated_cases["matching_case"])
    follow_up_case = cast(PatientExamination, annotated_cases["follow_up_case"])
    follow_up_report = cast(
        PatientExaminationReport,
        annotated_cases["follow_up_report"],
    )
    assert cases[1] == (
        "patient-pseudo-1",
        "center-a",
        f"{matching_case.pk}; {follow_up_case.pk}",
        "'=case-formula-prefix; case-follow-up",
        "2026-08-01; 2026-09-01",
        "colonoscopy; gastroscopy",
        "screening; surveillance",
        "gastritis; polyp",
        str(follow_up_report.pk),
        None,
        None,
    )
    assert len(cases) == 2
    assert metadata["Pseudonymous"] == "Yes"
    assert metadata["Row count"] == 1
    assert metadata["Grouping"] == "patient"
    assert "Sensitive" not in response.content.decode("latin-1")
    assert "Identity" not in response.content.decode("latin-1")
    assert annotated_cases["matching_case"] is not None


def test_xlsx_can_export_one_row_per_matching_examination(
    annotated_cases: dict[str, object],
) -> None:
    matching_case = cast(PatientExamination, annotated_cases["matching_case"])
    response = Client().get(
        "/endoreg-api/media/studies/case-export.xlsx",
        {
            "finding": "polyp",
            "indication": "screening",
            "group_by": "examination",
        },
    )

    assert response.status_code == 200
    assert response["X-Export-Row-Count"] == "1"
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    cases = list(workbook["Cases"].iter_rows(values_only=True))
    metadata = dict(workbook["Export metadata"].iter_rows(values_only=True))

    assert cases[0] == (
        "Case pseudonym",
        "Patient pseudonym",
        "Center",
        "Patient examination ID",
        "Examination date",
        "Examination",
        "Indications",
        "Findings",
        "Report IDs",
        "Processed report file IDs",
        "Processed video IDs",
    )
    assert cases[1] == (
        "'=case-formula-prefix",
        "patient-pseudo-1",
        "center-a",
        str(matching_case.pk),
        "2026-08-01",
        "colonoscopy",
        "screening",
        "polyp",
        None,
        None,
        None,
    )
    assert metadata["Grouping"] == "examination"


def test_xlsx_exports_the_exact_reviewed_study_cohort_snapshot(
    annotated_cases: dict[str, object],
) -> None:
    # Arrange
    matching_case = cast(PatientExamination, annotated_cases["matching_case"])
    payload = {
        "mode": "cohort",
        "group_by": "patient",
        "study_name": "Polyp registry 2026",
        "hypothesis": "Polyps are more frequent in the selected cohort.",
        "cohort_schema_version": "1.0",
        "patient_examination_id": [matching_case.pk],
        "date_from": "2026-01-01",
        "date_to": "2026-12-31",
        "center_key": "center-a",
        "examination_name": "colonoscopy",
        "finding": "polyp",
        "has_report": True,
        "has_video": False,
        "limit": 100,
    }

    # Act
    response = Client().post(
        "/endoreg-api/media/studies/case-export.xlsx",
        payload,
        content_type="application/json",
        HTTP_ACCEPT="application/json",
    )

    # Assert
    assert response.status_code == 200
    assert response["X-Export-Row-Count"] == "1"
    assert "pseudonymous-study-cohort-" in response["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    cases = list(workbook["Cases"].iter_rows(values_only=True))
    metadata = dict(workbook["Export metadata"].iter_rows(values_only=True))
    assert len(cases) == 2
    assert cases[1][2] == str(matching_case.pk)
    assert cases[1][5] == "colonoscopy"
    assert cases[1][8] is None
    assert metadata["Export mode"] == "cohort"
    assert metadata["Study name"] == "Polyp registry 2026"
    assert metadata["Hypothesis"] == (
        "Polyps are more frequent in the selected cohort."
    )
    assert metadata["Center"] == "center-a"
    assert metadata["Has report"] == "True"
    assert metadata["Has video"] == "False"
    assert metadata["Filter semantics"] == "Exact reviewed StudyCohortPage snapshot"


@pytest.mark.parametrize(
    ("query", "message"),
    [
        (
            {"mode": "cohort", "group_by": "patient"},
            "require study_name and hypothesis",
        ),
        (
            {
                "mode": "cohort",
                "group_by": "examination",
                "study_name": "Register",
                "hypothesis": "Hypothesis",
                "patient_examination_id": "1",
            },
            "must use patient grouping",
        ),
        (
            {
                "mode": "cohort",
                "group_by": "patient",
                "study_name": "Register",
                "hypothesis": "Hypothesis",
                "patient_examination_id": "999999",
            },
            "unavailable patient examination",
        ),
    ],
)
def test_xlsx_rejects_invalid_study_cohort_snapshots(
    annotated_cases: dict[str, object],
    query: dict[str, str],
    message: str,
) -> None:
    # Arrange
    del annotated_cases

    # Act
    response = Client().post(
        "/endoreg-api/media/studies/case-export.xlsx",
        query,
        content_type="application/json",
        HTTP_ACCEPT="application/json",
    )

    # Assert
    assert response.status_code == 400
    assert message in response.json()["error"]


@pytest.mark.parametrize(
    ("query", "message"),
    [
        ({}, "Select at least one"),
        ({"finding": "not-a-real-finding"}, "Unknown finding"),
        ({"finding": "polyp", "group_by": "invalid"}, "group_by must be"),
    ],
)
def test_xlsx_rejects_missing_or_unknown_selection(
    annotated_cases: dict[str, object],
    query: dict[str, str],
    message: str,
) -> None:
    del annotated_cases
    response = Client().get(
        "/endoreg-api/media/studies/case-export.xlsx",
        query,
    )

    assert response.status_code == 400
    assert message in response.json()["error"]


def test_xlsx_reports_empty_and_over_limit_selections(
    annotated_cases: dict[str, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del annotated_cases
    client = Client()
    empty_response = client.get(
        "/endoreg-api/media/studies/case-export.xlsx",
        {"examination": "gastroscopy", "indication": "screening"},
    )
    assert empty_response.status_code == 404
    assert "No pseudonymous cases" in empty_response.json()["error"]

    monkeypatch.setattr("lx_annotate.views.study_export.MAX_STUDY_EXPORT_ROWS", 0)
    over_limit_response = client.get(
        "/endoreg-api/media/studies/case-export.xlsx",
        {"finding": "polyp"},
    )
    assert over_limit_response.status_code == 400
    assert "narrow the selection" in over_limit_response.json()["error"]


def test_export_permission_requires_patient_read_outside_debug(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("lx_annotate.views.study_export.is_debug_mode", lambda: False)
    permission = StudyExportPatientReadPermission()

    class Groups:
        def __init__(self, roles: list[str]) -> None:
            self.roles = roles

        def values_list(self, field_name: str, flat: bool) -> list[str]:
            assert field_name == "name"
            assert flat is True
            return self.roles

    def request(*, authenticated: bool, roles: list[str]) -> Request:
        value = SimpleNamespace(
            user=SimpleNamespace(
                is_authenticated=authenticated,
                groups=Groups(roles),
            )
        )
        return cast(Request, cast(Any, value))

    assert (
        permission.has_permission(request(authenticated=False, roles=[]), object())
        is False
    )
    assert (
        permission.has_permission(
            request(authenticated=True, roles=["video:read"]), object()
        )
        is False
    )
    assert (
        permission.has_permission(
            request(authenticated=True, roles=["patient:read"]), object()
        )
        is True
    )
